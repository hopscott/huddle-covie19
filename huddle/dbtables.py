# pip install mysql-connector-python-rf
# https://stackoverflow.com/questions/10757169/location-of-my-cnf-file-on-macos
import mysql.connector
from mysql.connector import Error
from openpyxl import Workbook


# DB connection -----------------------------------------------------------
try:
    connection = mysql.connector.connect(host='localhost',
                                         database='icu_bed_covid19',
                                         user='root',
                                         password='huddleCOVIE19',
                                         auth_plugin='mysql_native_password')
    if connection.is_connected():
        db_Info = connection.get_server_info()
        print("Connected to MySQL Server version ", db_Info)
        cursor = connection.cursor()
        cursor.execute("select database();")
        record = cursor.fetchone()
        print("You're connected to database: ", record)

        # Table query ------------------------------------------------------
        SQL = "SELECT table_name FROM information_schema.tables WHERE table_schema = 'icu_bed_covid19';"
        cursor.execute(SQL)
        tableNames = cursor.fetchall()
        print(tableNames)

        # Create Excel (.xlsx) file ----------------------------------------
        wb = Workbook()

        for i in range(len(tableNames)):
            rawTable = tableNames[i]
            table = str(rawTable)
            print(table)
            SQL = 'SELECT * from ' + table + ';'
            cursor.execute(SQL)
            results = cursor.fetchall()
            ws = wb.create_sheet(i)
            ws.title = table
            ws.append(cursor.column_names)
            for row in results:
                ws.append(row)

        workbook_name = "test_workbook"
        wb.save(workbook_name + ".xlsx")

except Error as e:
    print("Error while connecting to MySQL", e)

finally:
    if (connection.is_connected()):
        cursor.close()
        connection.close()
        print("MySQL connection is closed")
